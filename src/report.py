"""5단계: 정답지 대조(탐지율), 검토결과 저장, 결과 엑셀 생성 (CLAUDE.md 2-4, 3, 5-2)."""
from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path

import pandas as pd

ANSWER_SHEET = "정답지"
REVIEW_CHOICES = ["", "적정", "소명요청", "부적정"]
NOTICE = ("※ AI 선별 결과는 점검 필요 거래를 고르기 위한 참고자료이며, "
          "부적정 여부의 최종 판단은 감사부 담당자가 증빙 확인 후 내립니다.")
# 앱에 아직 넣지 않은 규칙 (CLAUDE.md 2-2의 Nice 항목)
NOT_IMPLEMENTED = {"R11", "R12", "R14", "R16", "R17"}

# 결과 엑셀 '점검대상목록' 열 (CLAUDE.md 3장 점검결과 출력 양식)
EXPORT_COLUMNS = {
    "승인번호": "승인번호", "승인일시": "승인일시", "직원명": "직원", "예산부서명": "부서", "가맹점명": "가맹점",
    "전표금액": "금액", "걸린 규칙": "걸린 규칙", "위험점수": "위험점수", "등급": "등급",
    "AI 선별 사유": "AI 사유", "AI 추가 확인사항": "추가확인사항", "검토결과": "검토결과", "검토의견": "검토의견",
}


# ── 검토결과 저장 (CSV 1개, 승인번호/패턴키 기준) ─────────────────


def load_reviews(path: Path) -> pd.DataFrame:
    if path.exists():
        r = pd.read_csv(path, dtype=str).fillna("")
        return r
    return pd.DataFrame(columns=["키", "검토결과", "검토의견", "수정일시"])


def save_reviews(path: Path, current: pd.DataFrame, edited: pd.DataFrame) -> pd.DataFrame:
    """편집한 행(키·검토결과·검토의견)을 기존 저장분에 덮어써 저장하고, 저장된 전체를 돌려준다."""
    edited = edited[["키", "검토결과", "검토의견"]].fillna("").astype(str)
    old = current.set_index("키") if len(current) else pd.DataFrame(columns=["검토결과", "검토의견", "수정일시"])
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    for _, r in edited.iterrows():
        prev = old.loc[r["키"]] if r["키"] in old.index else None
        changed = prev is None or (prev["검토결과"], prev["검토의견"]) != (r["검토결과"], r["검토의견"])
        if not changed:
            continue
        if not r["검토결과"] and not r["검토의견"]:
            old = old.drop(index=r["키"], errors="ignore")
        else:
            old.loc[r["키"]] = [r["검토결과"], r["검토의견"], now]
    out = old.reset_index().rename(columns={"index": "키"})
    path.parent.mkdir(parents=True, exist_ok=True)
    out.to_csv(path, index=False, encoding="utf-8-sig")
    return out


def attach_reviews(df: pd.DataFrame, key: pd.Series, reviews: pd.DataFrame) -> pd.DataFrame:
    r = reviews.set_index("키") if len(reviews) else pd.DataFrame(columns=["검토결과", "검토의견"])
    out = df.copy()
    out["키"] = key.astype(str).values
    out["검토결과"] = out["키"].map(r["검토결과"]).fillna("")
    out["검토의견"] = out["키"].map(r["검토의견"]).fillna("")
    return out


# ── 정답지 대조 ─────────────────────────────────────────────


def read_answers(source) -> pd.DataFrame:
    a = pd.read_excel(source, sheet_name=ANSWER_SHEET, dtype={"승인번호": "int64"})
    # '— 규칙(R07)만으로는 못 잡음' 같은 설명 뒷부분의 규칙ID는 기대 규칙이 아니므로 '—' 앞만 본다
    a["기대규칙"] = a["기대 탐지규칙"].fillna("").map(
        lambda s: sorted(set(re.findall(r"T-R\d|R\d\d", s.split("—")[0]))))
    return a


def evaluate(answers: pd.DataFrame, scored: pd.DataFrame, selected: pd.DataFrame,
             patterns: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """시나리오별 탐지 여부와 요약 지표를 계산한다. 정답지는 채점에만 쓰고 앱 판단에는 쓰지 않는다."""
    sel_ids = set(selected["승인번호"])
    s_idx = scored.set_index("승인번호")
    pat_ids: dict[str, set] = {}
    for _, p in patterns.iterrows():
        code = p["패턴"].split()[0]
        pat_ids.setdefault(code, set()).update(int(i) for i in p["관련 승인번호"].split(", "))

    rows = []
    for sid, g in answers.groupby("시나리오ID", sort=False):
        ids = list(g["승인번호"])
        expected = g["기대규칙"].iloc[0]
        if sid == "FP":
            for _, r in g.iterrows():
                x = s_idx.loc[r["승인번호"]]
                ok = x["등급"] in ("하", "") and x["AI 조정"] < 0
                rows.append(_row("FP", [r["승인번호"]], r, ok, f"{x['등급']} {x['위험점수']}점 (조정 {x['AI 조정']})",
                                 "저위험 조정됨" if ok else "조정 안 됨", expected, []))
            continue
        if sid.startswith("T"):
            code = next((c for c in expected if c.startswith("T-R")), "")
            hit_ids = pat_ids.get(code, set())
            ok = any(i in hit_ids for i in ids)
            rows.append(_row(sid, ids, g.iloc[0], ok, f"{sum(i in hit_ids for i in ids)}/{len(ids)}건 패턴 목록 포함",
                             "탐지" if ok else "미탐", expected, [code] if ok else []))
            continue
        hit_rules = set()
        for i in ids:
            hit_rules.update(re.findall(r"R\d\d", s_idx.loc[i, "걸린 규칙"]))
        ok = all(i in sel_ids for i in ids)
        missing = [r for r in expected if r not in hit_rules]
        note = "탐지" if ok else ("미탐 — 미구현 규칙(" + ", ".join(sorted(set(expected) & NOT_IMPLEMENTED)) + ") 필요"
                                if set(expected) <= NOT_IMPLEMENTED else "미탐")
        grade = s_idx.loc[ids[0], "등급"]
        rows.append(_row(sid, ids, g.iloc[0], ok, f"{grade} {s_idx.loc[ids[0], '위험점수']}점", note, expected,
                         [r for r in expected if r in hit_rules], missing))

    result = pd.DataFrame(rows)
    s, t, fp = (result[result["시나리오"].str.startswith(p)] for p in ("S", "T", "FP"))
    exp_rules = sum(len(x) for x in s["기대 규칙"].str.split(", ") if x != [""])
    got_rules = sum(len(x) for x in s["맞힌 규칙"].str.split(", ") if x != [""])
    summary = {
        "단건 시나리오(S) 탐지": f"{s['탐지'].sum()}/{len(s)} ({s['탐지'].mean():.0%})",
        "기간 패턴(T) 탐지": f"{t['탐지'].sum()}/{len(t)} ({t['탐지'].mean():.0%})",
        "오탐(FP) 저위험 조정": f"{fp['탐지'].sum()}/{len(fp)} ({fp['탐지'].mean():.0%})",
        "S 기대 규칙 적중": f"{got_rules}/{exp_rules} ({got_rules / exp_rules:.0%})",
    }
    return result, summary


def _row(sid, ids, r, ok, result, note, expected, got, missing=None):
    return {
        "시나리오": sid, "승인번호": ", ".join(map(str, ids)), "가맹점": r["가맹점명"], "시나리오 설명": r["시나리오"],
        "기대 규칙": ", ".join(expected), "맞힌 규칙": ", ".join(got),
        "놓친 규칙": ", ".join(missing or []), "앱 결과": result, "탐지": bool(ok), "판정": note,
    }


# ── 결과 엑셀 ──────────────────────────────────────────────


def build_excel(month: str, selected: pd.DataFrame, patterns: pd.DataFrame, recs: pd.DataFrame,
                eval_df: pd.DataFrame | None, eval_summary: dict | None) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    target = selected[list(EXPORT_COLUMNS)].rename(columns=EXPORT_COLUMNS).copy()
    target.insert(0, "No", range(1, len(target) + 1))
    target["승인일시"] = target["승인일시"].dt.strftime("%Y-%m-%d %H:%M")
    grades = selected["등급"].value_counts()
    head = (f"법인카드 점검대상 목록 — {month}  |  선별 {len(target)}건 (상 {grades.get('상', 0)} · "
            f"중 {grades.get('중', 0)} · 하 {grades.get('하', 0)})")

    pat = patterns.drop(columns=["_가중"], errors="ignore").copy()
    sheets = [(f"점검대상목록_{month}", head, target),
              ("기간패턴분석", f"기간(시계열) 패턴 분석 — 기준월 {month}", pat),
              ("AI추천", "AI 추천 — 기존 점검기준 외 추가 검토 항목", recs)]
    if eval_df is not None:
        sm = "  |  ".join(f"{k} {v}" for k, v in (eval_summary or {}).items())
        sheets.append(("정답지대조", f"정답지 대조 — {sm}", eval_df))

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, title, data in sheets:
            data.to_excel(xw, sheet_name=name, index=False, startrow=3)
            ws = xw.sheets[name]
            ws["A1"] = title
            ws["A1"].font = Font(bold=True, size=13)
            ws["A2"] = NOTICE
            ws["A2"].font = Font(color="666666", size=9)
            for cell in ws[4]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E8EEF7")
            for i, col in enumerate(data.columns, start=1):
                longest = max([len(str(col))] + [len(str(v)) for v in data[col].head(200)])
                ws.column_dimensions[get_column_letter(i)].width = min(max(8, longest * 1.6), 60)
                for row in ws.iter_rows(min_row=5, min_col=i, max_col=i):
                    row[0].alignment = Alignment(wrap_text=longest > 30, vertical="top")
                    if col in ("금액", "전표금액"):
                        row[0].number_format = "#,##0"
            ws.freeze_panes = "A5"
    return buf.getvalue()
