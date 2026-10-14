"""가상 법인카드 거래 1개월분 생성기.

기존 가상데이터 엑셀의 '전체' 시트(SAP 추출 34개 컬럼)를 참고해, 같은 컬럼·같은 형식으로
지정한 달의 정상 거래를 만든다. 이상거래는 심지 않는다.

  python -I tools/make_month_dummy.py data/법인카드_가상데이터_202509-202608.xlsx 2026-09 1000 \
      data/법인카드_가상데이터_202609.xlsx --holidays 2026-09-24,2026-09-25,2026-09-26,2026-09-28

만드는 방식
- 원본의 '정상' 거래(정답지·취소·휴일 행 제외)를 한 건씩 골라 카드·직원·부서·계정·업종·적요 조합을 그대로 쓰고,
  승인일(해당 월 영업일)·승인 시각(원본 시간대 유지)·금액(원본 ±25%)만 새로 정한다.
- 거래량이 원본보다 많아도 한 가맹점에 몰리지 않도록 업종·지역별 가상 가맹점을 추가한다.
- 결재문서번호·전표번호는 카드×월마다 1건, 전표행번호는 승인일시 순으로 매긴다(원본 규칙과 같음).
- 승인번호·전표번호·결재문서번호·가맹점사업자번호는 원본의 마지막 번호 다음부터 이어 붙인다.
- 취소 거래는 원본 비율(약 1.3%)로 매입·취소 쌍을 만든다.
"""
from __future__ import annotations

import argparse
import calendar
import datetime as dt
import random

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

TEXT_COLS = ["G/L 계정", "승인시간", "예산부서", "업종코드", "가맹점사업자번호", "전표행번호", "회계연도"]

# 거래량이 늘어도 한 가맹점에 몰리지 않게 추가하는 가상 가맹점 (업종, 지역(시·도 시·군·구), 이름)
EXTRA_MERCHANTS = [
    ("일반음식점", "서울 중구", ["샘플한식당", "가상순두부", "더미칼국수", "예시설렁탕", "모의돈까스", "테스트쌈밥", "가상냉면", "샘플생선구이"]),
    ("일반음식점", "서울 강남구", ["모의한우", "예시파스타", "가상중화요리", "샘플초밥"]),
    ("일반음식점", "서울 종로구", ["더미한정식", "예시보쌈"]),
    ("일반음식점", "부산 부산진구", ["가상돼지국밥", "예시복국", "모의아구찜"]),
    ("일반음식점", "부산 해운대구", ["샘플장어", "테스트횟집"]),
    ("커피전문점", "서울 중구", ["예시빈스", "가상커피랩", "모의로스터리", "샘플카페"]),
    ("커피전문점", "서울 강남구", ["더미커피"]),
    ("커피전문점", "부산 부산진구", ["가상카페 서면점", "예시커피 부산점"]),
    ("제과점", "서울 중구", ["가상베이커리", "샘플제과"]),
    ("제과점", "부산 부산진구", ["예시빵집 부산점"]),
    ("편의점", "서울 중구", ["가상마트 을지로점", "샘플24 명동점"]),
    ("편의점", "부산 부산진구", ["예시25 서면점"]),
    ("꽃집", "서울 중구", ["가상플라워", "모의꽃집"]),
    ("주차장", "서울 중구", ["예시파킹 을지로점", "가상주차장 명동점"]),
    ("주차장", "부산 부산진구", ["모의파킹 서면점"]),
    ("주유소", "서울 중구", ["샘플에너지 주유소", "더미오일 중구주유소"]),
    ("백화점", "서울 중구", ["모의백화점 명동점"]),
    ("백화점", "서울 강남구", ["가상백화점 강남점", "예시아울렛 강남점"]),
    ("주유소", "부산 부산진구", ["가상오일 부산주유소"]),
    ("문구/사무용품", "서울 중구", ["더미문구"]),
    ("PG일반(인증)", "서울 강남구", ["가상오픈마켓"]),
]
NO_DISTANCE = {"PG일반(인증)", "해외가맹점", "택시", "항공사", "철도", "주유소", "주차장"}  # R14 제외 업종
STREETS = ["예시로", "가상길", "샘플대로", "더미길", "모의대로", "테스트로", "예시대로", "가상로"]


def load_source(path):
    df = pd.read_excel(path, sheet_name="전체", dtype={c: str for c in TEXT_COLS})
    cols = list(df.columns)
    try:
        ans = pd.read_excel(path, sheet_name="정답지")
        planted = set(ans["승인번호"].astype(int))
    except ValueError:
        planted = set()
    normal = df[(df["취소구분"] == "매입") & (df["상태"].isna()) & (~df["승인번호"].isin(planted))
                & (df["국내해외코드명"] == "국내")]
    return df, cols, normal


def business_days(year, month, holidays):
    days = []
    for d in range(1, calendar.monthrange(year, month)[1] + 1):
        day = dt.date(year, month, d)
        if day.weekday() < 5 and day not in holidays:
            days.append(day)
    return days


def amount_like(src_amt, lo, hi, limit, rng):
    """원본 금액을 ±25% 흔들되, 원본의 단위(1,000/100/10원)와 업종 범위를 유지한다."""
    unit = 1000 if src_amt % 1000 == 0 else 100 if src_amt % 100 == 0 else 10
    for _ in range(20):
        x = round(src_amt * rng.uniform(0.75, 1.25) / unit) * unit
        x = min(max(x, lo), hi)
        if limit:
            x = min(x, limit)
        if x >= 100000 and x % 100000 == 0:  # 10만원 단위 정액은 원본에서도 드물어서 피함
            continue
        return int(x)
    return int(src_amt)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("month", help="YYYY-MM")
    ap.add_argument("count", type=int)
    ap.add_argument("output")
    ap.add_argument("--holidays", default="", help="쉼표로 구분한 공휴일 YYYY-MM-DD")
    ap.add_argument("--seed", type=int, default=20260930)
    a = ap.parse_args()

    rng = random.Random(a.seed)
    year, month = map(int, a.month.split("-"))
    holidays = {dt.date.fromisoformat(x) for x in a.holidays.split(",") if x}
    days = business_days(year, month, holidays)
    month_end = dt.date(year, month, calendar.monthrange(year, month)[1])

    src, cols, normal = load_source(a.source)
    limits = pd.read_excel(a.source, sheet_name="계정코드", dtype=str).set_index("G/L 계정")["건당 한도(원)"].astype(int).to_dict()
    pos = normal[normal["전표금액"] > 0]
    ranges = pos.groupby("가맹점업종")["전표금액"].agg(["min", "max"]).to_dict("index")

    # 가맹점 풀: 원본 가맹점 + 추가 가상 가맹점 (업종, 시·도별)
    merchants = src.drop_duplicates("가맹점명")[["가맹점업종", "가맹점명", "가맹점주소", "가맹점사업자번호", "업종코드"]].to_dict("records")
    code_of = {m["가맹점업종"]: m["업종코드"] for m in merchants}
    next_biz = int(src["가맹점사업자번호"].max()) + 1
    for ind, area, names in EXTRA_MERCHANTS:
        for name in names:
            merchants.append({"가맹점업종": ind, "가맹점명": name, "가맹점주소": f"{area} {rng.choice(STREETS)} {rng.randint(3, 299)}",
                              "가맹점사업자번호": str(next_biz), "업종코드": code_of[ind]})
            next_biz += 1
    pool = {}
    for m in merchants:
        pool.setdefault((m["가맹점업종"], m["가맹점주소"].split()[0]), []).append(m)

    # 1) 정상 거래 뽑기
    n_cancel = round(a.count * 25 / 1857)
    n_purchase = a.count - n_cancel
    rows = []
    for _ in range(n_purchase):
        s = pos.sample(1, random_state=rng.randint(0, 2**31 - 1)).iloc[0].to_dict()
        r = dict(s)
        cands = pool.get((s["가맹점업종"], s["가맹점주소"].split()[0]), [])
        if cands and rng.random() < 0.6:
            m = rng.choice(cands)
            r.update({k: m[k] for k in ["가맹점명", "가맹점주소", "가맹점사업자번호", "업종코드"]})
        day = rng.choice(days)
        hh = int(s["승인시간"][:2])
        r["_dt"] = dt.datetime(day.year, day.month, day.day, hh, rng.randint(0, 59), rng.randint(0, 59))
        rg = ranges[s["가맹점업종"]]
        r["전표금액"] = amount_like(int(s["전표금액"]), rg["min"], rg["max"], limits.get(s["G/L 계정"]), rng)
        rows.append(r)

    # 1-1) 같은 카드가 짧은 간격으로 다른 시·도에서 결제한 것처럼 보이면(R14 우연 적중) 날짜를 다시 뽑는다
    for _ in range(200):
        moved = 0
        by_card = {}
        for r in rows:
            by_card.setdefault(r["카드번호"], []).append(r)
        for g in by_card.values():
            g.sort(key=lambda r: r["_dt"])
            for x, y in zip(g, g[1:]):
                if x["가맹점업종"] in NO_DISTANCE or y["가맹점업종"] in NO_DISTANCE:
                    continue
                if x["가맹점주소"].split()[0] != y["가맹점주소"].split()[0] and y["_dt"] - x["_dt"] <= dt.timedelta(minutes=120):
                    day = rng.choice(days)
                    y["_dt"] = y["_dt"].replace(year=day.year, month=day.month, day=day.day)
                    moved += 1
        if not moved:
            break

    # 2) 취소 쌍: 매입 직후 몇 분 뒤 같은 금액을 취소
    for p in rng.sample(rows, n_cancel):
        c = dict(p)
        c["_dt"] = (p["_dt"] + dt.timedelta(minutes=rng.randint(2, 9))).replace(second=0)
        c["전표금액"] = -p["전표금액"]
        c["취소구분"] = "취소"
        c["전표 품목 텍스트"] = f"{p['전표 품목 텍스트']} (승인취소)"
        rows.append(c)

    # 3) 번호 매기기: 승인번호(승인일시 순), 카드×월 결재문서·전표, 전표행번호
    rows.sort(key=lambda r: r["_dt"])
    appr = int(src["승인번호"].max())
    voucher = int(src["전표번호"].max())
    doc_seq = int(src["결재문서번호"].str[4:16].astype(int).max())
    per_card = {}
    for r in rows:
        appr += rng.randint(3, 37)
        r["승인번호"] = appr
        card = r["카드번호"]
        if card not in per_card:
            voucher += rng.randint(3, 12)
            doc_seq += 1
            if r["구분"] == "공용법인":
                title = f"{year % 100}년 {month}월 부서공용카드"
            else:
                title = rng.choice([f"{month}월 경비사용내역", f"{year % 100}.{month:02d} 법인카드 사용내역"])
            per_card[card] = {"전표번호": voucher, "결재문서번호": f"FI22{doc_seq:012d}{year}", "결재문서제목": title, "line": 0}
        v = per_card[card]
        v["line"] += 1
        amt = r["전표금액"]
        supply = round(amt / 1.1)
        r.update({
            "상태": None, "승인일": int(r["_dt"].strftime("%Y%m%d")), "승인시간": r["_dt"].strftime("%H%M%S"),
            "전표번호": v["전표번호"], "결재문서번호": v["결재문서번호"], "결재문서제목": v["결재문서제목"], "전표행번호": str(v["line"]),
            "공급가액": supply, "승인부가세": amt - supply, "승인봉사료": 0, "외화거래일환율": 0.0, "통화": "KRW", "국내해외코드명": "국내",
            "회계일자": dt.datetime.combine(month_end, dt.time()), "회계연도": str(year),
        })

    write(a.output, cols, rows)
    print(f"{a.output}: {len(rows)}건 (매입 {n_purchase}, 취소 {n_cancel}), 카드 {len(per_card)}장, 승인번호 {rows[0]['승인번호']}~{rows[-1]['승인번호']}")


def write(path, cols, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "전체"
    ws.append(cols)
    head_fill = PatternFill("solid", fgColor="DDE6F0")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([None if (isinstance(v := r.get(c), float) and v != v) else v for c in cols])
    money = {"전표금액", "공급가액", "승인부가세", "승인봉사료"}
    for j, name in enumerate(cols, start=1):
        letter = openpyxl.utils.get_column_letter(j)
        if name in money:
            for (cell,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                cell.number_format = "#,##0"
        if name == "회계일자":
            for (cell,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                cell.number_format = "yyyy-mm-dd"
        width = max(len(str(name)) * 2, *(len(str(r.get(name) or "")) * 1.3 for r in rows[:200])) + 2
        ws.column_dimensions[letter].width = min(max(width, 8), 40)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


if __name__ == "__main__":
    main()
